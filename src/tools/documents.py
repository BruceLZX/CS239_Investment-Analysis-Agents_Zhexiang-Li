"""Read a client document as text with location markers, so an analyst can cite what it read as "<file> p.N".

Each format, and the marker its blocks of text get:
- .md, .txt, .csv: the text as it is, marked by line ranges only when it runs over one page of output;
- .pdf: "p.N", a block a page;
- .docx: "section N: <heading>", a block from each heading to the next, with tables, then text boxes, headers and
  footers, and comments;
- .pptx: "slide N", a block a slide, with tables, charts, alt text and speaker notes;
- .xlsx: "sheet <name>", a line a row, starting with the row number;
- scanned PDF pages and images (.png, .jpg, .tif and others): read by OCR when the pytesseract package and the
  tesseract program are both installed, and tagged OCR_TAG. Without OCR, a document with no text at all comes back
  as a result starting with NEEDS_OCR, which the caller turns into a request for a text version.

Nothing in the file is dropped because a reader of the page wouldn't see it: speaker notes, hidden slides, sheets and
rows, hidden or tiny text, off-slide shapes, comments and document properties are kept and labelled, so a scan for
instructions aimed at AI tools (rule N8) sees them. Characters that show as nothing but can split a word (zero-width
spaces, direction marks) are removed and counted, and Unicode tag characters, which spell out hidden ASCII text, are
decoded. PDF text drawn in white is extracted like any other text; pypdf doesn't report colour.

A page of output holds at most about MAX_CHARS characters: read_document(path, page=2) returns the second, and every
page but the last ends by naming the next.
"""
from dataclasses import dataclass, field
from datetime import date, datetime, time
import functools
from itertools import groupby
from pathlib import Path
import re
import shutil
import unicodedata

MAX_CHARS = 9_000
OCR_TAG = '[OCR, low confidence]'
NEEDS_OCR = '[needs OCR or a text version]'
TEXT_KINDS = {'.md': 'Markdown', '.markdown': 'Markdown', '.txt': 'text', '.csv': 'CSV', '.tsv': 'TSV'}
IMAGE_SUFFIXES = ('.png', '.jpg', '.jpeg', '.tif', '.tiff', '.bmp', '.gif', '.webp')
LEGACY = {'.doc': '.docx', '.ppt': '.pptx', '.xls': '.xlsx'}
# Characters that show as nothing but split a word, so that "ign\u200bore" slips past a scan for "ignore".
INVISIBLE_RE = re.compile('[\u00ad\u061c\u180e\u200b-\u200f\u202a-\u202e\u2060-\u2064\u2066-\u2069\ufeff]')
# Unicode tag characters spell out ASCII text that screens don't show but language models read.
TAGS_RE = re.compile('[\U000e0000-\U000e007f]+')
# A PDF page with an image and fewer non-space characters than this is a scan.
MIN_PAGE_TEXT = 20
# Text set smaller than this many points can't be read on the page.
TINY_POINTS = 4
NO_TEXT_LAYER = '(an image with no text layer: it needs OCR or a text version)'
SHEET_STATES = {'hidden': '[hidden sheet]', 'veryHidden': '[very hidden sheet: only a macro can show it]'}


@dataclass
class Block:
    location: str  # 'p.3', 'slide 4', 'sheet Revenue', 'section 2: Market'; '' for a whole text file
    text: str


@dataclass
class Reading:
    kind: str  # e.g. 'PDF, 15 pages'
    cite: str  # how a location is cited after the file name, e.g. 'p.N'
    blocks: list[Block]
    needs_ocr: list[str] = field(default_factory=list)  # locations with no text layer, read without OCR
    ocr: bool = False  # some text came from OCR
    invisible: int = 0  # invisible characters removed


def read_document(path: str | Path, page: int = 1) -> str:
    """One page of the document's text (at most about MAX_CHARS characters) with its location markers, a header
    naming the file, its format and how to cite it, and a closing line naming the next page."""
    path = Path(path)
    reading = _reading(path)
    if _unreadable(reading):
        return _needs_ocr(path, reading)
    pages = _paginate(reading.blocks)
    if not 1 <= page <= len(pages):
        raise ValueError(f'{path.name} has {_plural(len(pages), "page")} of output; ask for page 1 to {len(pages)}.')
    cite = f'{path.name} {reading.cite}'.strip()
    lines = [f'{path.name}: {reading.kind}. Output page {page} of {len(pages)}{_span(pages[page - 1])}. '
             f'Cite as "{cite}".'] + [f'Note: {note}' for note in _notes(reading)]
    body = '\n\n'.join(_rendered(label, text) for label, text in pages[page - 1])
    if page < len(pages):
        end = f'--- More: output page {page + 1} of {len(pages)} has{_span(pages[page])}; read it with page={page + 1}. ---'
    else:
        end = f'--- End of {path.name}. ---'
    return '\n'.join(lines) + '\n\n' + body + '\n\n' + end


def document_text(path: str | Path) -> str:
    """The whole document with its location markers and notes, unpaginated, for a scan over all of it."""
    path = Path(path)
    reading = _reading(path)
    if _unreadable(reading):
        return _needs_ocr(path, reading)
    return '\n\n'.join([f'Note: {note}' for note in _notes(reading)] +
                       [_rendered(block.location, block.text) for block in reading.blocks])


def _reading(path: Path) -> Reading:
    stat = path.stat()
    return _cached_reading(str(path.resolve()), stat.st_mtime_ns, stat.st_size, _ocr_engine() is not None)


@functools.lru_cache(maxsize=32)
def _cached_reading(path: str, mtime_ns: int, size: int, ocr: bool) -> Reading:
    """The file read once per version (and OCR setting), so paging through it doesn't parse or OCR it again."""
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix in TEXT_KINDS:
        reading = _text(path)
    elif suffix == '.pdf':
        reading = _pdf(path)
    elif suffix == '.docx':
        reading = _docx(path)
    elif suffix == '.pptx':
        reading = _pptx(path)
    elif suffix in ('.xlsx', '.xlsm'):
        reading = _xlsx(path)
    elif suffix in IMAGE_SUFFIXES:
        reading = _image(path)
    elif suffix in LEGACY:
        raise ValueError(f'{path.name} is in the old binary {suffix} format, which this reader can\'t open; ask the '
                         f'client for {LEGACY[suffix]} or PDF.')
    else:
        raise ValueError(f'{path.name}: {suffix or "files without an extension"} can\'t be read here. Readable: '
                         + ', '.join([*TEXT_KINDS, '.pdf', '.docx', '.pptx', '.xlsx', *IMAGE_SUFFIXES]) + '.')
    for block in reading.blocks:
        block.text, removed = _clean(block.text)
        reading.invisible += removed
    return reading


def _clean(text: str) -> tuple[str, int]:
    """Text a scan can match, and how many invisible characters were removed from it: tag characters decoded,
    compatibility forms (ligatures, full-width letters, non-breaking spaces) folded and blank runs collapsed."""
    text = TAGS_RE.sub(lambda m: '[hidden tag characters: ' + ''.join(
        chr(ord(c) - 0xE0000) for c in m.group() if 0xE0020 <= ord(c) < 0xE007F) + ']', text)
    text, removed = INVISIBLE_RE.subn('', text)
    text = unicodedata.normalize('NFKC', text).replace('\r\n', '\n').replace('\r', '\n')
    text = '\n'.join(line.rstrip() for line in text.split('\n'))
    return re.sub(r'\n{3,}', '\n\n', text).strip('\n'), removed


def _unreadable(reading: Reading) -> bool:
    """Every part with content needs OCR, so there is nothing to read without it."""
    skip = set(reading.needs_ocr) | {'document properties'}
    return bool(reading.needs_ocr) and not any(b.text.strip() for b in reading.blocks if b.location not in skip)


def _needs_ocr(path: Path, reading: Reading) -> str:
    return (f'{NEEDS_OCR} {path.name} ({reading.kind}) has no text layer: it is a scan or an image, and OCR isn\'t '
            'available here (it needs the pytesseract package and the tesseract program). Ask the client for a text '
            'version, such as the original PDF, Word or PowerPoint file.')


def _notes(reading: Reading) -> list[str]:
    notes = []
    if reading.needs_ocr:
        verbs = ('has', 'needs') if len(reading.needs_ocr) == 1 else ('have', 'need')
        notes.append(f'{", ".join(reading.needs_ocr)} {verbs[0]} no text layer and {verbs[1]} OCR or a text version.')
    if reading.ocr:
        notes.append(f'text tagged {OCR_TAG} was read by OCR and may hold misread figures; check them against the '
                     'original.')
    if reading.invisible:
        notes.append(f'removed {_plural(reading.invisible, "invisible character")} (zero-width spaces, direction '
                     'marks or soft hyphens), which can hide words from a reader or a scan.')
    return notes


def _plural(count: int, noun: str) -> str:
    return f'{count} {noun}{"" if count == 1 else "s"}'


def _rendered(label: str, text: str) -> str:
    return f'--- {label} ---\n{text}' if label else text


def _span(items: list[tuple[str, str]]) -> str:
    labels = [label.split(':')[0].removesuffix(' (continued)') for label, _ in items if label]
    if not labels:
        return ''
    return f' ({labels[0]})' if labels[0] == labels[-1] else f' ({labels[0]} to {labels[-1]})'


def _paginate(blocks: list[Block]) -> list[list[tuple[str, str]]]:
    """Blocks packed into pages of at most about MAX_CHARS characters; a block too long for one page is split."""
    pages, current, size = [], [], 0
    for block in blocks:
        for label, text in _pieces(block, MAX_CHARS - 200):
            cost = len(label) + len(text) + 10
            if current and size + cost > MAX_CHARS:
                pages.append(current)
                current, size = [], 0
            current.append((label, text))
            size += cost
    if current:
        pages.append(current)
    return pages or [[('', '(no text)')]]


def _pieces(block: Block, limit: int) -> list[tuple[str, str]]:
    """The block as (label, text) pieces of at most limit characters, cut at line breaks where it can be. Pieces of
    a whole text file are labelled by line range; later pieces of any other block are labelled continued."""
    if len(block.text) <= limit:
        return [(block.location, block.text)]
    lines = [(number, line[start:start + limit]) for number, line in enumerate(block.text.split('\n'), 1)
             for start in range(0, max(len(line), 1), limit)]
    chunks, current, size = [], [], 0
    for number, line in lines:
        if current and size + len(line) + 1 > limit:
            chunks.append(current)
            current, size = [], 0
        current.append((number, line))
        size += len(line) + 1
    chunks.append(current)
    if block.location:
        labels = [block.location] + [f'{block.location} (continued)'] * (len(chunks) - 1)
    else:
        labels = [f'lines {chunk[0][0]}-{chunk[-1][0]}' for chunk in chunks]
    return [(label, '\n'.join(line for _, line in chunk)) for label, chunk in zip(labels, chunks)]


def _ocr_engine():
    """pytesseract's image_to_string when the package and the tesseract program are both installed, else None."""
    try:
        import pytesseract
    except ImportError:
        return None
    return pytesseract.image_to_string if shutil.which(pytesseract.pytesseract.tesseract_cmd) else None


def _properties(title, subject, keywords, comments) -> list[Block]:
    """The document's own title, subject, keywords and comments, where instructions can hide too."""
    fields = [f'{name}: {" ".join(str(value).split())}' for name, value in
              (('title', title), ('subject', subject), ('keywords', keywords), ('comments', comments))
              if value and str(value).strip()]
    return [Block('document properties', '\n'.join(fields))] if fields else []


def _labelled(label: str, text: str) -> str:
    return f'[{label}: {text}]' if label and text.strip() else text


def _text(path: Path) -> Reading:
    raw = path.read_bytes()
    try:
        text = raw.decode('utf-8-sig')
    except UnicodeDecodeError:
        text = raw.decode('cp1252', errors='replace')
    lines = len(text.rstrip('\n').split('\n'))
    return Reading(f'{TEXT_KINDS[path.suffix.lower()]}, {_plural(lines, "line")}', '', [Block('', text)])


def _pdf(path: Path) -> Reading:
    from pypdf import PdfReader

    reader = PdfReader(path)
    if reader.is_encrypted and not reader.decrypt(''):
        raise ValueError(f'{path.name} is password-protected; ask the client for an unlocked copy.')
    info = reader.metadata
    reading = Reading(f'PDF, {_plural(len(reader.pages), "page")}', 'p.N',
                      _properties(info.title, info.subject, info.keywords, None) if info else [])
    engine = _ocr_engine()
    for number, page in enumerate(reader.pages, 1):
        location, text = f'p.{number}', page.extract_text() or ''
        images = _page_images(page) if len(re.sub(r'\s+', '', text)) < MIN_PAGE_TEXT else []
        if not images:
            reading.blocks.append(Block(location, text if text.strip() else '(no text on this page)'))
        elif engine:
            reading.ocr = True
            read = [engine(image) if image is not None else '(an image OCR could not decode)' for image in images]
            reading.blocks.append(Block(location, '\n'.join(filter(str.strip, [text, OCR_TAG, *read]))))
        else:
            reading.needs_ocr.append(location)
            reading.blocks.append(Block(location, '\n'.join(filter(str.strip, [text, NO_TEXT_LAYER]))))
    return reading


def _page_images(page) -> list:
    """The page's images, decoded for OCR; None for one that pypdf or Pillow can't decode."""
    try:
        files = list(page.images)
    except Exception:  # A broken or unsupported image dictionary: treat the page as having an image.
        return [None]
    images = []
    for file in files:
        try:
            images.append(file.image)
        except Exception:  # A filter Pillow can't decode, such as JBIG2.
            images.append(None)
    return images


def _docx(path: Path) -> Reading:
    import docx
    from docx.text.paragraph import Paragraph

    document = docx.Document(path)
    props = document.core_properties
    blocks = _properties(props.title, props.subject, props.keywords, props.comments)
    location, lines, headings = 'before the first heading', [], 0
    for item in document.iter_inner_content():
        if not isinstance(item, Paragraph):
            lines += _docx_table(item)
            continue
        text = _docx_paragraph(item)
        style = (item.style.name if item.style is not None else '').lower()
        if text.strip() and style.startswith(('heading', 'title')):
            if lines:
                blocks.append(Block(location, '\n'.join(lines)))
            headings += 1
            location, lines = f'section {headings}: {" ".join(text.split())[:80]}', []
        if text.strip():
            lines.append(text)
    if lines:
        blocks.append(Block(location if headings else 'whole document', '\n'.join(lines)))
    blocks += [Block(name, text) for name, text in (('text boxes', _docx_text_boxes(document)),
                                                     ('headers and footers', _docx_headers(document)),
                                                     ('comments', _docx_comments(document))) if text]
    return Reading(f'Word, {_plural(headings, "heading")}', 'section N', blocks)


def _docx_paragraph(paragraph) -> str:
    runs = [run for item in paragraph.iter_inner_content() for run in getattr(item, 'runs', [item])]
    return ''.join(_labelled(label, ''.join(run.text for run in group)) for label, group in groupby(runs, _docx_hidden))


def _docx_hidden(run) -> str:
    """Why a reader of the page wouldn't see the run, or '' if they would."""
    font = run.font
    if font.hidden:
        return 'hidden text'
    if font.size is not None and font.size.pt < TINY_POINTS:
        return 'tiny text'
    if font.color.type is not None and str(font.color.rgb) == 'FFFFFF':
        return 'white text'
    return ''


def _docx_table(table) -> list[str]:
    rows = []
    for row in table.rows:
        cells, last = [], None
        for cell in row.cells:
            if cell._tc is not last:  # A merged cell repeats across the columns it spans.
                cells.append(' '.join(cell.text.split()))
            last = cell._tc
        while cells and not cells[-1]:
            cells.pop()
        if cells:
            rows.append(' | '.join(cells))
    return rows


def _docx_text_boxes(document) -> str:
    from docx.oxml.ns import qn

    boxes = []
    for box in document.element.body.iter(qn('w:txbxContent')):
        text = '\n'.join(''.join(t.text or '' for t in p.iter(qn('w:t'))) for p in box.iter(qn('w:p'))).strip()
        if text and text not in boxes:  # Word saves each text box twice, as DrawingML and as a VML fallback.
            boxes.append(text)
    return '\n'.join(boxes)


def _docx_headers(document) -> str:
    parts = []
    for section in document.sections:
        for part in (section.header, section.footer):
            if part.is_linked_to_previous:
                continue
            text = '\n'.join(p.text for p in part.paragraphs if p.text.strip())
            if text and text not in parts:
                parts.append(text)
    return '\n'.join(parts)


def _docx_comments(document) -> str:
    return '\n'.join(f'[comment] {" ".join(comment.text.split())}' for comment in document.comments
                     if comment.text.strip())


def _pptx(path: Path) -> Reading:
    from pptx import Presentation

    presentation = Presentation(path)
    props = presentation.core_properties
    blocks = _properties(props.title, props.subject, props.keywords, props.comments)
    size = (presentation.slide_width, presentation.slide_height)
    slides = list(presentation.slides)
    for number, slide in enumerate(slides, 1):
        lines = ['[hidden slide: skipped in the slideshow]'] if slide._element.get('show') in ('0', 'false') else []
        for shape in _reading_order(slide.shapes):
            lines += _pptx_shape(shape, size)
        notes = slide.notes_slide.notes_text_frame if slide.has_notes_slide else None
        if notes is not None and notes.text.strip():
            lines.append('[speaker notes] ' + notes.text.strip())
        blocks.append(Block(f'slide {number}', '\n'.join(lines) or '(no text on this slide)'))
    return Reading(f'PowerPoint, {_plural(len(slides), "slide")}', 'slide N', blocks)


def _reading_order(shapes) -> list:
    return sorted(shapes, key=lambda s: (s.top is None, s.top or 0, s.left or 0))


def _pptx_shape(shape, size: tuple[int, int] | None) -> list[str]:
    """The shape's text, tables, chart data and alt text, labelled when the shape sits outside the slide."""
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    lines = []
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        for child in _reading_order(shape.shapes):
            lines += _pptx_shape(child, None)  # A child's position is in the group's own coordinates.
    if shape.has_text_frame:
        for paragraph in shape.text_frame.paragraphs:
            if any(_pptx_hidden(run) for run in paragraph.runs):
                text = ''.join(_labelled(label, ''.join(r.text for r in group))
                               for label, group in groupby(paragraph.runs, _pptx_hidden))
            else:
                text = paragraph.text.replace('\v', '\n')
            if text.strip():
                lines.append(text)
    if shape.has_table:
        for row in shape.table.rows:
            cells = [' '.join(cell.text.split()) for cell in row.cells]
            if any(cells):
                lines.append(' | '.join(cells))
    if shape.has_chart:
        lines += _chart(shape.chart)
    alt = [text.strip() for text in shape._element.xpath('./*[1]/p:cNvPr/@descr') if text.strip()]
    lines += [f'[alt text] {text}' for text in alt]
    if lines and size and _off_slide(shape, size):
        lines.insert(0, '[off-slide: placed outside the visible slide]')
    return lines


def _pptx_hidden(run) -> str:
    return 'tiny text' if run.font.size is not None and run.font.size.pt < TINY_POINTS else ''


def _off_slide(shape, size: tuple[int, int]) -> bool:
    width, height = size
    if None in (shape.left, shape.top, shape.width, shape.height):
        return False
    return (shape.left >= width or shape.top >= height or shape.left + shape.width <= 0
            or shape.top + shape.height <= 0)


def _chart(chart) -> list[str]:
    """The chart's title, categories and series values: decks often give their figures only as charts."""
    try:
        title = chart.chart_title.text_frame.text if chart.has_title and chart.chart_title.has_text_frame else ''
        lines = [f'[chart] {" ".join(title.split())}'.rstrip()]
        for plot in chart.plots:
            categories = [str(category) for category in plot.categories]
            if categories:
                lines.append('categories: ' + ' | '.join(categories))
            for series in plot.series:
                lines.append(f'{series.name}: ' + ' | '.join('' if v is None else _number(v, '') for v in series.values))
        return lines
    except Exception:  # A chart type python-pptx can't parse: say so rather than fail the whole deck.
        return ['[chart: its data could not be read]']


def _xlsx(path: Path) -> Reading:
    import openpyxl

    values = openpyxl.load_workbook(path, data_only=True)  # Each formula's result as last saved.
    formulas = openpyxl.load_workbook(path)  # The formulas, for cells saved without a result.
    props = values.properties
    blocks = _properties(props.title, props.subject, props.keywords, props.description)
    for sheet in values.worksheets:
        lines = [SHEET_STATES[sheet.sheet_state]] if sheet.sheet_state in SHEET_STATES else []
        hidden = [letter for letter, dimension in sheet.column_dimensions.items() if dimension.hidden]
        if hidden:
            lines.append('[hidden columns: ' + ', '.join(hidden) + ']')
        comments = []
        for row in sheet.iter_rows():
            cells = [_cell(cell, formulas[sheet.title][cell.coordinate]) for cell in row]
            while cells and not cells[-1]:
                cells.pop()
            number = row[0].row
            if cells:
                flag = ' [hidden row]' if sheet.row_dimensions[number].hidden else ''
                lines.append(f'{number}{flag}: ' + ' | '.join(cells))
            comments += [f'[comment on {cell.coordinate}] {" ".join(cell.comment.text.split())}'
                         for cell in row if getattr(cell, 'comment', None)]
        blocks.append(Block(f'sheet {sheet.title}', '\n'.join(lines + comments) or '(empty sheet)'))
    return Reading(f'Excel, {_plural(len(values.worksheets), "sheet")}', 'sheet <name>, row N', blocks)


def _cell(cell, formula_cell) -> str:
    value = cell.value
    if value is None:
        formula = formula_cell.value
        return str(formula) if isinstance(formula, str) and formula.startswith('=') else ''
    if isinstance(value, bool):
        return 'TRUE' if value else 'FALSE'
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == time(0) else value.isoformat(sep=' ', timespec='minutes')
    if isinstance(value, (date, time)):
        return value.isoformat()
    if isinstance(value, (int, float)):
        return _number(value, cell.number_format or '')
    return ' '.join(str(value).split())


def _number(value, number_format: str) -> str:
    if '%' in number_format:
        return f'{value * 100:.15g}%'
    return str(value) if isinstance(value, int) else f'{value:.15g}'


def _image(path: Path) -> Reading:
    from PIL import Image, ImageSequence

    engine = _ocr_engine()
    with Image.open(path) as image:
        count = getattr(image, 'n_frames', 1)
        frames = [frame.copy() for frame in ImageSequence.Iterator(image)] if engine else []
    locations = [f'p.{n}' for n in range(1, count + 1)] if count > 1 else ['image']
    kind = f'image, {_plural(count, "page")}'
    cite = 'p.N' if count > 1 else ''
    if not engine:
        return Reading(kind, cite, [Block(location, NO_TEXT_LAYER) for location in locations], needs_ocr=locations)
    return Reading(kind, cite, [Block(location, f'{OCR_TAG}\n{engine(frame)}')
                                for location, frame in zip(locations, frames)], ocr=True)
